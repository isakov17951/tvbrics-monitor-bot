# -*- coding: utf-8 -*-
"""
Экспорт архива найденных новостей в Excel, сгруппировано по партнёрам.
Запускается вручную на GitHub (Actions -> "Экспорт в Excel" -> Run workflow),
там же задаются временные рамки. Готовый файл приходит в Telegram.

Переменные окружения:
  BOT_TOKEN, CHAT_ID — как у поиска
  DATE_FROM, DATE_TO — рамки в формате ДД.ММ.ГГГГ (пустые = без ограничения)
"""

import asyncio
import os
import sqlite3
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import DB_PATH

TOKEN = os.getenv("BOT_TOKEN")
CHAT_ID = os.getenv("CHAT_ID", "")
CHAT_IDS = [c.strip() for c in CHAT_ID.split(",") if c.strip()]
DATE_FROM = (os.getenv("DATE_FROM") or "").strip()
DATE_TO = (os.getenv("DATE_TO") or "").strip()

FONT = "Arial"


def parse_dmy(s):
    return datetime.strptime(s, "%d.%m.%Y") if s else None


def load_rows():
    d_from, d_to = parse_dmy(DATE_FROM), parse_dmy(DATE_TO)
    conn = sqlite3.connect(DB_PATH)
    rows = conn.execute(
        "SELECT partner, news_date, title, url, lang FROM seen WHERE partner != ''"
    ).fetchall()
    conn.close()
    out = []
    for partner, news_date, title, url, lang in rows:
        try:
            d = datetime.strptime(news_date, "%d.%m.%Y")
        except (ValueError, TypeError):
            continue
        if d_from and d < d_from:
            continue
        if d_to and d > d_to:
            continue
        # в одной записи может быть "IRNA, Mehr News" — раскладываем по обоим
        for p in [p.strip() for p in partner.split(",")]:
            out.append((p, d, title, url, lang))
    return out


def build_xlsx(rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Новости партнёров"

    title_font = Font(name=FONT, size=14, bold=True)
    partner_font = Font(name=FONT, size=12, bold=True, color="FFFFFF")
    partner_fill = PatternFill("solid", start_color="2F5597")
    head_font = Font(name=FONT, size=10, bold=True)
    head_fill = PatternFill("solid", start_color="D9E2F3")
    cell_font = Font(name=FONT, size=10)
    link_font = Font(name=FONT, size=10, color="0563C1", underline="single")

    period = (f"с {DATE_FROM or 'начала архива'} по {DATE_TO or 'сегодня'}")
    ws["A1"] = f"Новости иранских партнёров TV BRICS ({period})"
    ws["A1"].font = title_font

    r = 3
    by_partner = {}
    for p, d, title, url, lang in rows:
        by_partner.setdefault(p, []).append((d, title, url, lang))

    for partner in sorted(by_partner):
        items = sorted(by_partner[partner], key=lambda x: x[0])
        ws.cell(row=r, column=1, value=f"{partner} — {len(items)} шт.")
        for col in range(1, 5):
            ws.cell(row=r, column=col).fill = partner_fill
        ws.cell(row=r, column=1).font = partner_font
        r += 1
        for col, name in enumerate(["Дата", "Заголовок", "Ссылка", "Язык"], 1):
            c = ws.cell(row=r, column=col, value=name)
            c.font, c.fill = head_font, head_fill
        r += 1
        for d, title, url, lang in items:
            c = ws.cell(row=r, column=1, value=d.strftime("%d.%m.%Y"))
            c.font = cell_font
            c = ws.cell(row=r, column=2, value=title)
            c.font = cell_font
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c = ws.cell(row=r, column=3, value=url)
            c.hyperlink, c.font = url, link_font
            c = ws.cell(row=r, column=4, value=lang)
            c.font = cell_font
            c.alignment = Alignment(horizontal="center")
            r += 1
        r += 1  # пустая строка между партнёрами

    for col, width in zip("ABCD", (12, 55, 60, 8)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    wb.save(path)


async def send(path, caption):
    from aiogram import Bot
    from aiogram.types import FSInputFile
    bot = Bot(token=TOKEN)
    try:
        for chat_id in CHAT_IDS:
            try:
                await bot.send_document(chat_id, FSInputFile(path), caption=caption)
            except Exception as e:
                print(f"Не смог отправить получателю {chat_id}: {e}", flush=True)
    finally:
        await bot.session.close()


def main():
    if not os.path.exists(DB_PATH):
        raise SystemExit("База seen_news.db не найдена: архив пуст, "
                         "сначала должен отработать хотя бы один поиск.")
    rows = load_rows()
    print(f"Записей в выгрузке: {len(rows)}", flush=True)
    fname = "iran_news_{}_{}.xlsx".format(
        (DATE_FROM or "start").replace(".", "-"),
        (DATE_TO or datetime.now().strftime("%d-%m-%Y")).replace(".", "-"),
    )
    build_xlsx(rows, fname)
    caption = (f"Экспорт: {len(rows)} ссылок, "
               f"{DATE_FROM or 'начало архива'} — {DATE_TO or 'сегодня'}")
    asyncio.run(send(fname, caption))


if __name__ == "__main__":
    main()
